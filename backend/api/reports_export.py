from decimal import Decimal
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse

NAVY = '1E3A8A'
LIGHT_GRAY = 'F1F5F9'
HEADER_FILL = PatternFill(start_color=NAVY, end_color=NAVY, fill_type='solid')
HEADER_FONT = Font(color='FFFFFF', bold=True)
TITLE_FONT = Font(size=15, bold=True, color=NAVY)
SUBTITLE_FONT = Font(size=10, color='475569')
META_FONT = Font(size=10, bold=True, color='1E293B')
THIN_BORDER = Border(bottom=Side(style='thin', color='CBD5E1'))
GS_FORMAT = '#,##0 "Gs."'


def build_xlsx_response(filename, sheet_title, columns, rows, report_title=None, extra_meta=None):
    """
    columns: lista de dicts {'header': str, 'width': int, 'format': 'gs'|'text'|'number'|None}
    rows: lista de tuplas/listas con los valores en el mismo orden que columns
    report_title: título visible del reporte (ej. "Clientes con mora"). Si no se
        pasa, se usa sheet_title.
    extra_meta: lista opcional de strings con líneas de contexto adicionales
        (ej. "Período: 01/06/2026 a 17/09/2026", "Filtro: categoría Heladeras").
    """
    from .models import CompanyInfo

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    company = CompanyInfo.objects.first()
    total_cols = max(len(columns), 4)

    row_cursor = 1

    ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=total_cols)
    ws.cell(row=row_cursor, column=1, value=company.name if company else 'Tienda').font = TITLE_FONT
    row_cursor += 1

    if company:
        contact_bits = [b for b in [
            f'RUC: {company.ruc}' if company.ruc else None,
            f'Tel: {company.phone}' if company.phone else None,
            company.address or None,
        ] if b]
        if contact_bits:
            ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=total_cols)
            ws.cell(row=row_cursor, column=1, value=' · '.join(contact_bits)).font = SUBTITLE_FONT
            row_cursor += 1

    row_cursor += 1

    ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=total_cols)
    ws.cell(row=row_cursor, column=1, value=report_title or sheet_title).font = Font(size=12, bold=True)
    row_cursor += 1

    now = timezone.now()
    meta_lines = [
        f'Generado: {now.strftime("%d/%m/%Y %H:%M")}',
        f'Total de registros: {len(rows)}',
    ] + (extra_meta or [])
    for line in meta_lines:
        ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=total_cols)
        ws.cell(row=row_cursor, column=1, value=line).font = META_FONT
        row_cursor += 1

    row_cursor += 1
    header_row = row_cursor

    for col_idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=col['header'])
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[get_column_letter(col_idx)].width = col.get('width', 18)

    for row_offset, row in enumerate(rows, start=1):
        row_idx = header_row + row_offset
        for col_idx, (col, value) in enumerate(zip(columns, row), start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = THIN_BORDER
            if col.get('format') == 'gs':
                cell.number_format = GS_FORMAT
            elif col.get('format') == 'number':
                cell.number_format = '#,##0'

    ws.freeze_panes = f'A{header_row + 1}'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def to_number(value):
    if value is None:
        return 0
    if isinstance(value, Decimal):
        return float(value)
    return value
